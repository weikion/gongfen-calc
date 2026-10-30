"""excel_io —— 文件读取、Sheet 选择、表头定位、合并单元格展开。

职责边界：本模块只负责"把 Excel 变成规整矩阵 + 元信息"，
不做字段识别、不做单位换算、不做打分。
"""
from __future__ import annotations

import csv
import difflib
import io
import re
import unicodedata
from pathlib import Path

import params as P

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    openpyxl = None


# ============================================================ 通用文本工具
def norm_str(v) -> str:
    """全角→半角、去不间断空格、统一中文标点。所有文本比较前的唯一入口。"""
    if v is None:
        return ""
    s = unicodedata.normalize("NFKC", str(v))
    s = s.replace("\u3000", " ").replace("\xa0", " ")
    s = s.replace("，", ",").replace("；", ";").replace("：", ":")
    s = s.replace("（", "(").replace("）", ")")
    return s.strip()


def is_null(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        s = norm_str(v).lower()
        return s in P.NULL_TOKENS
    return False


def to_float(v):
    """把可能的数值/文本数值转 float，失败返回 None。不处理单位（交给 value_parse）。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = norm_str(v).replace(",", "").replace(" ", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def similarity(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a, b).ratio() * 100.0


# ============================================================ 文件读取
class ParseError(Exception):
    """统一解析异常。code 必须是 E_* 枚举值，echo 为给用户看的可回显内容。"""

    def __init__(self, code: str, msg: str, echo=None):
        super().__init__(f"[{code}] {msg}")
        self.code = code
        self.msg = msg
        self.echo = echo or {}


ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "utf-16")
GARBLED_MARKS = ("锟斤拷", "烫烫烫")


def _decode_bytes(raw: bytes):
    for enc in ENCODINGS:
        try:
            text = raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        bad = text.count("\ufffd")
        if bad / max(len(text), 1) < 0.001 and not any(m in text for m in GARBLED_MARKS):
            return text, enc
    return None, None


def _matrix_from_csv(path: Path):
    raw = path.read_bytes()
    text, enc = _decode_bytes(raw)
    if text is None:
        raise ParseError(
            "E_ENCODING",
            "编码无法识别（尝试过 utf-8-sig/utf-8/gb18030/gbk/big5/utf-16）",
            {"hint": "请用 Excel 另存为「CSV UTF-8」或提供 xlsx 文件"},
        )
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        sep = dialect.delimiter
    except csv.Error:
        sep = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=sep))
    matrix = [[c if c != "" else None for c in row] for row in rows]
    return matrix, enc


def _sheet_meta(ws, matrix):
    """提取合并区、隐藏行列、列数字格式。"""
    merges = []
    for rng in ws.merged_cells.ranges:
        merges.append([rng.min_row, rng.min_col, rng.max_row, rng.max_col])

    hidden_rows = [r for r, dim in ws.row_dimensions.items() if dim is not None and dim.hidden]
    hidden_cols = []
    for col_letter, dim in ws.column_dimensions.items():
        if dim is not None and dim.hidden:
            try:
                hidden_cols.append(openpyxl.utils.column_index_from_string(col_letter))
            except Exception:
                pass

    col_formats = {}
    sample = min(len(matrix), 60)
    n_cols = max((len(r) for r in matrix), default=0)
    for c in range(1, n_cols + 1):
        fmts = set()
        for r in range(1, sample + 1):
            try:
                fmt = ws.cell(r, c).number_format
            except Exception:
                continue
            if fmt:
                fmts.add(fmt)
        if fmts:
            col_formats[c] = sorted(fmts)

    return {
        "name": ws.title,
        "matrix": matrix,
        "merges": merges,
        "hidden_rows": hidden_rows,
        "hidden_cols": hidden_cols,
        "col_formats": col_formats,
        "n_rows": len(matrix),
        "n_cols": n_cols,
    }


def load(path) -> dict:
    """读文件 → bundle{path, format, sheets[], formula_uncomputed, encoding}。"""
    path = Path(path)
    if not path.exists():
        raise ParseError("E_FILE_NOT_FOUND", f"文件不存在：{path}", {"path": str(path)})

    suffix = path.suffix.lower()
    if suffix == ".csv":
        matrix, enc = _matrix_from_csv(path)
        return {
            "path": str(path),
            "format": "csv",
            "encoding": enc,
            "formula_uncomputed": 0,
            "sheets": [{
                "name": path.stem, "matrix": matrix, "merges": [],
                "hidden_rows": [], "hidden_cols": [], "col_formats": {},
                "n_rows": len(matrix),
                "n_cols": max((len(r) for r in matrix), default=0),
            }],
        }

    if suffix in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        if openpyxl is None:
            raise ParseError(
                "E_DEP_MISSING",
                "缺少 openpyxl",
                {"hint": "pip install -r requirements.txt -i https://mirrors.aliyun.com/pypi/simple/ --trusted-host mirrors.aliyun.com"},
            )
        try:
            wb_v = openpyxl.load_workbook(path, data_only=True)
            wb_f = openpyxl.load_workbook(path, data_only=False)
        except Exception as exc:  # 加密 / 损坏
            raise ParseError(
                "E_FILE_UNREADABLE",
                f"文件无法打开（可能加密或损坏）：{type(exc).__name__}",
                {"hint": "若文件有打开密码，请先解密另存后重传", "detail": str(exc)[:200]},
            ) from exc

        sheets = []
        uncomputed = 0
        for ws in wb_v.worksheets:
            max_row, max_col = ws.max_row, ws.max_column
            matrix = [[None] * max_col for _ in range(max_row)]
            for r in range(1, max_row + 1):
                for c in range(1, max_col + 1):
                    matrix[r - 1][c - 1] = ws.cell(r, c).value
            # 公式未计算检测
            ws_f = wb_f[ws.title]
            for r in range(1, min(max_row, 200) + 1):
                for c in range(1, max_col + 1):
                    fv = ws_f.cell(r, c).value
                    if isinstance(fv, str) and fv.startswith("=") and ws.cell(r, c).value is None:
                        uncomputed += 1
            meta = _sheet_meta(ws, matrix)
            meta["has_autofilter"] = ws.auto_filter.ref is not None
            meta["autofilter_first_row"] = (
                int("".join(ch for ch in ws.auto_filter.ref.split(":")[0] if ch.isdigit()))
                if ws.auto_filter.ref else None
            )
            sheets.append(meta)

        return {
            "path": str(path),
            "format": "xlsx",
            "encoding": None,
            "formula_uncomputed": uncomputed,
            "sheets": sheets,
        }

    if suffix == ".xls":
        try:
            import pandas as pd
            df_map = pd.read_excel(path, sheet_name=None, header=None)
        except Exception as exc:
            raise ParseError(
                "E_FILE_UNREADABLE",
                "旧版 .xls 读取失败（需 xlrd>=2.0）",
                {"hint": "建议用 Excel 另存为 .xlsx 后重传", "detail": str(exc)[:200]},
            ) from exc
        sheets = []
        for name, df in df_map.items():
            matrix = df.where(df.notna(), None).values.tolist()
            sheets.append({
                "name": str(name), "matrix": matrix, "merges": [],
                "hidden_rows": [], "hidden_cols": [], "col_formats": {},
                "n_rows": len(matrix), "n_cols": max((len(r) for r in matrix), default=0),
                "has_autofilter": False, "autofilter_first_row": None,
            })
        return {"path": str(path), "format": "xls", "encoding": None,
                "formula_uncomputed": 0, "sheets": sheets}

    raise ParseError(
        "E_FORMAT_UNSUPPORTED",
        f"不支持的文件类型：{suffix}",
        {"hint": "支持 .xlsx / .xlsm / .xls / .csv"},
    )


# ============================================================ Sheet 评分选择
DOC_SHEET_WORDS = re.compile(r"说明|汇总|统计|填表|模板|readme|目录|封面|口径|字段|定义", re.I)


def _doc_sheet_penalty(sheet) -> float:
    matrix = sheet["matrix"]
    name = str(sheet["name"])
    if DOC_SHEET_WORDS.search(name):
        return 1.0
    if sheet["n_rows"] < 5:
        return 1.0
    # 单列超宽且全为长句 → 说明页
    if sheet["n_cols"] <= 2:
        longs = [v for row in matrix[:40] for v in row
                 if isinstance(v, str) and len(v) > 40]
        if len(longs) >= 5:
            return 1.0
    return 0.0


def pick_sheet(bundle, probe_row) -> dict:
    """按 alias 命中 / 填充 / 数据量 / autofilter 综合评分选表。

    probe_row(row_values) -> int  (该行命中的别名个数)
    分差过小 → E_SHEET_AMBIGUOUS（绝不猜）。
    """
    scored = []
    for sheet in bundle["sheets"]:
        matrix = sheet["matrix"]
        if not matrix:
            continue
        alias_hits = 0
        data_rows = 0
        for r, row in enumerate(matrix[:P.HEADER_SCAN_ROWS]):
            hits = probe_row(row)
            if hits > alias_hits:
                alias_hits = hits
            if r > 0 and sum(0 if is_null(v) else 1 for v in row) >= 2:
                data_rows += 1
        score = (
            4.0 * min(alias_hits / 3.0, 1.0)
            + 3.0 * min(alias_hits / 4.0, 1.0)
            + 2.0 * min(data_rows / 50.0, 1.0)
            + 3.0 * (1.0 if sheet.get("has_autofilter") else 0.0)
            - 5.0 * _doc_sheet_penalty(sheet)
        )
        scored.append({"name": sheet["name"], "score": round(score, 3),
                       "alias_hits": alias_hits, "n_rows": sheet["n_rows"]})

    if not scored:
        raise ParseError("E_EMPTY_FILE", "文件中没有任何可读的 sheet", {})

    scored.sort(key=lambda x: -x["score"])
    best = scored[0]
    candidates = [s for s in scored if s["alias_hits"] > 0]

    if not candidates:
        preview = []
        for sh in bundle["sheets"]:
            m = sh["matrix"]
            cols = []
            for c in range(min(sh["n_cols"], 15)):
                vals = [norm_str(m[r][c]) for r in range(min(len(m), 20))
                        if c < len(m[r]) and not is_null(m[r][c])][:3]
                cols.append({"col_idx": c + 1, "sample_values": vals})
            preview.append({"name": sh["name"], "n_rows": sh["n_rows"], "columns": cols})
        raise ParseError(
            "E_NO_LIST_SHEET",
            "未找到像任务清单的 sheet（无任何列名命中别名词典）",
            {"sheets": scored, "preview": preview,
             "hint": "表头可能不在前 15 行，或列名全为非标准命名。请按下列列样例手动映射"},
        )

    if len(candidates) > 1 and (candidates[0]["score"] - candidates[1]["score"]) < P.SHEET_SCORE_GAP:
        raise ParseError(
            "E_SHEET_AMBIGUOUS",
            f"发现 {len(candidates)} 个候选清单表，分差不足 {P.SHEET_SCORE_GAP}，需用户确认",
            {"candidates": candidates, "hint": "请指定使用哪个 sheet"},
        )

    chosen_name = best["name"] if best["alias_hits"] > 0 else candidates[0]["name"]
    sheet = next(s for s in bundle["sheets"] if s["name"] == chosen_name)
    return {"sheet": sheet, "candidates": scored}


# ============================================================ 合并单元格展开
def apply_merges(sheet):
    """展开合并单元格。

    返回 (matrix, fill_map)
    fill_map[(r0, c0)] = {"anchor": (ar0, ac0), "span_rows": n, "span_cols": m}
    纵向合并**不直接当作重复任务**，由下游按 anchor 分组聚合。
    """
    matrix = [list(row) for row in sheet["matrix"]]
    n_cols = sheet["n_cols"]
    for row in matrix:
        if len(row) < n_cols:
            row.extend([None] * (n_cols - len(row)))

    fill_map = {}
    for min_row, min_col, max_row, max_col in sheet["merges"]:
        ar, ac = min_row - 1, min_col - 1
        anchor_val = matrix[ar][ac]
        span_rows, span_cols = max_row - min_row + 1, max_col - min_col + 1
        for r in range(ar, max_row):
            for c in range(ac, max_col):
                matrix[r][c] = anchor_val
                fill_map[(r, c)] = {"anchor": (ar, ac), "span_rows": span_rows,
                                    "span_cols": span_cols}
    return matrix, fill_map


# ============================================================ 表头定位
UNIT_ROW_RE = re.compile(
    r"^\(?单位[:：]?.*\)?$|^\(.{1,6}\)$|^(元|万元|小时|人天|人日|天|%|个|条|人|次)$"
)
HEADER_PENALTY_RE = re.compile(r"合计|总计|^单位[:：]|^注[:：]|说明|备注|制表|编制|审核")


def _data_row_score(row) -> float:
    if row is None:
        return 0.0
    nonempty = [v for v in row if not is_null(v)]
    if not nonempty:
        return 0.0
    fill = len(nonempty) / max(len(row), 1)
    numeric = sum(1 for v in nonempty if isinstance(v, (int, float))) / len(nonempty)
    long_text = any(isinstance(v, str) and len(v) >= 4 for v in nonempty)
    return 0.6 * fill + 0.4 * max(numeric, 1.0 if long_text else 0.0)


def detect_header(sheet, probe_row) -> dict:
    """前 15 行打分定位表头；识别双层表头与单位行。"""
    matrix = sheet["matrix"]
    n_cols = max(sheet["n_cols"], 1)
    best = None

    for r in range(min(P.HEADER_SCAN_ROWS, len(matrix))):
        row = matrix[r]
        nonempty = [v for v in row if not is_null(v)]
        if not nonempty:
            continue
        fill = len(nonempty) / n_cols
        texts = [norm_str(v) for v in nonempty if isinstance(v, str)]
        short_ratio = (sum(1 for t in texts if len(t) <= 14) / len(nonempty)) if nonempty else 0.0
        hits = probe_row(row)
        h3 = hits / n_cols
        h4 = _data_row_score(matrix[r + 1]) if r + 1 < len(matrix) else 0.0
        # 惩罚项：仅在本行几乎不可能是表头时才生效（避免「备注」列名误伤真表头）
        pen_words = 1.0 if (hits <= 1 and any(
            HEADER_PENALTY_RE.search(norm_str(v)) for v in nonempty)) else 0.0
        pen_title = 1.0 if len(nonempty) <= 1 else 0.0
        # 表头行不应含数值：本行数值占比高 → 更像数据行
        nums_self = sum(1 for v in nonempty if isinstance(v, (int, float)))
        pen_data = 1.0 if (nums_self / max(len(nonempty), 1) >= 0.2) else 0.0
        bonus = 0.0
        if sheet.get("autofilter_first_row") == r + 1:
            bonus = 1.5
        score = (3.0 * h3 + 2.0 * fill + 1.5 * short_ratio
                 + 2.0 * h4 - 3.0 * pen_words - 2.5 * pen_title - 2.5 * pen_data + bonus)
        if best is None or score > best["score"]:
            best = {"row_index": r, "score": round(score, 3), "hits": hits,
                    "detected_by": ("autofilter" if bonus else "alias+shape")}

    if best is None or best["score"] < P.HEADER_SCORE_THRESHOLD:
        raise ParseError(
            "E_HEADER_NOT_FOUND",
            f"表头定位失败（前 {P.HEADER_SCAN_ROWS} 行最高分 "
            f"{best['score'] if best else 0:.2f} < 阈值 {P.HEADER_SCORE_THRESHOLD}）",
            {"head_rows": [
                [norm_str(v) for v in row[:20]]
                for row in matrix[:P.HEADER_SCAN_ROWS]
            ]},
        )

    r = best["row_index"]
    row_span = 1

    def _fill(i):
        if i < 0 or i >= len(matrix):
            return 0.0
        return sum(0 if is_null(v) else 1 for v in matrix[i]) / n_cols

    cur_fill = _fill(r)

    # 双层表头（向上找父行）：打分偏好「填得更满」的子行，所以父行常在 r-1。
    # 判据三件套：本行像表头(hits>=2) + 本行不满(fill<=0.75) + 存在纵向合并痕迹
    # （父有值、子为空），三者同时成立才升级，避免把标题/编制人行误判成父表头。
    if r > 0:
        up = matrix[r - 1]
        up_hits = probe_row(up)
        up_fill = _fill(r - 1)
        child = matrix[r]
        has_vmerge = any(
            (c < len(up) and not is_null(up[c]))
            and (c >= len(child) or is_null(child[c]))
            for c in range(n_cols)
        )
        if up_hits >= 2 and up_fill <= 0.75 and has_vmerge:
            r, row_span = r - 1, 2

    # 双层表头（向下找子行）：本行不满（有合并腾空）+ 下一行也像表头（命中 >=2、无任何数值）
    # → 判为两级表头。「本行不满」与「下一行无任何数值」是两个关键闸门，
    # 满行单层表头 + 首行数据（必然含数值）都不会被误判。
    if row_span == 1 and r + 1 < len(matrix):
        nxt = matrix[r + 1]
        nxt_fill = _fill(r + 1)
        nxt_hits = probe_row(nxt)
        nxt_nums = sum(1 for v in nxt if not is_null(v) and isinstance(v, (int, float)))
        if cur_fill <= 0.8 and nxt_hits >= 2 and nxt_fill >= 0.4 and nxt_nums == 0:
            row_span = 2

    unit_row = None
    unit_hint = None
    probe_unit = r + row_span
    if probe_unit < len(matrix):
        cells = [norm_str(v) for v in matrix[probe_unit] if not is_null(v)]
        if cells and all(UNIT_ROW_RE.match(c) for c in cells):
            unit_row = probe_unit
            for c in cells:
                m = re.search(r"单位[:：]?\s*([^)）]+)", c)
                if m:
                    unit_hint = m.group(1).strip()
                    break
            if unit_hint is None and len(cells) == 1:
                unit_hint = cells[0].strip("()（）")

    data_start = r + row_span + (1 if unit_row is not None else 0)
    return {
        "row_index": r, "row_span": row_span, "unit_row": unit_row,
        "unit_hint": unit_hint, "confidence": min(best["score"], 1.0),
        "detected_by": best["detected_by"], "data_start": data_start,
    }


# 非任务行关键词：个人考核表里这些是说明/意见栏，不是任务，必须剔除，
# 否则「自我月评」「下月计划」这类长文本会被当成任务并因含"优化/迁移"等词拿到高分。
NON_TASK_ROW_RE = re.compile(
    r"自我月评|自我总结|自我评价|下月计划|本月计划|上月总结|工作计划|季度计划|"
    r"考核意见|部门意见|分管|领导意见|意见|签字|审核人|批准人|制表人|"
    r"额外加分|其他/|其他／|分值合计|在职人员|岗位职责")


NUM_TOKEN_RE = re.compile(r"^[\d,.\s]+$")
# 多区块表头专用：只拦合计类整行（不拦「备注」「说明」这类合法列名）
BLOCK_HEADER_PENALTY_RE = re.compile(r"^合\s*计|^总\s*计|^累\s*计|^汇\s*总|^单位[:：]")


def is_non_task_row(row) -> bool:
    """区块标题/说明/意见类行：不是任务，不参与打分。"""
    nonempty = [norm_str(v) for v in row if not is_null(v)]
    if not nonempty:
        return True
    # 单格文本 = 区块标题（如「技术功能开发」）
    if len(nonempty) == 1:
        return True
    # 合并单元格展开后整行同值 = 区块标题（如「日常工作考核（自行更替内容）」）
    if len(set(nonempty)) == 1:
        return True
    return any(NON_TASK_ROW_RE.search(v) for v in nonempty)


def _looks_like_block_header(row, next_row) -> bool:
    """短文本表头行判据（多区块用）。

    与 detect_header 的打分不同：考核表每个区块的表头列数很少（2–4 列），
    别名命中数经常只有 1，靠 hits>=2 会漏判。改用形态判据。
    三条硬闸门（缺一不可，都是实测踩出来的）：
      ① 非空 >=2 且**值不全相同** —— 排除合并展开的区块标题行（整行同值）
      ② 无数值（含纯数字字符串）—— 排除数据行（序号列常是文本型数字）
      ③ 全部短文本（<=14 字）—— 排除含长描述的数据行
    """
    nonempty = [v for v in row if not is_null(v)]
    if len(nonempty) < 2:
        return False
    if len({norm_str(v) for v in nonempty}) == 1:
        return False
    for v in nonempty:
        if isinstance(v, (int, float)):
            return False
        if isinstance(v, str) and NUM_TOKEN_RE.match(v.strip()):
            return False
    texts = [norm_str(v) for v in nonempty]
    if any(len(t) > 14 for t in texts):
        return False
    # 只排除合计类整行，不能复用 HEADER_PENALTY_RE —— 后者含「备注|说明」，
    # 而「备注」恰恰是考核表里最常见的合法列名，照搬会把整个区块的表头判没。
    if any(BLOCK_HEADER_PENALTY_RE.search(t) for t in texts):
        return False
    return _data_row_score(next_row) > 0.3 if next_row is not None else False


def _is_block_title_row(row) -> bool:
    """区块标题行：整行同值（合并单元格展开）或只有一个非空格。
    如「技术功能开发」「日常工作考核（自行更替内容）」。"""
    nonempty = [norm_str(v) for v in row if not is_null(v)]
    if not nonempty:
        return False
    if len(nonempty) == 1:
        return len(nonempty[0]) >= 2 and not NUM_TOKEN_RE.match(nonempty[0])
    return len(set(nonempty)) == 1


def detect_blocks(matrix) -> list:
    """多区块表：找出全部表头行并切分数据区。

    依据：个人月度考核表常见「区块标题 → 独立表头 → 若干数据行」重复结构。
    单表头解析只能覆盖第一个（或打分最高的）区块，会漏掉其余区块的任务，
    并把月评/计划/意见栏当成任务计入。
    入参必须是**合并展开后**的矩阵，否则合并的区块标题行只有 1 个非空格、判不出来。

    关键闸门：第 2 个及以后的表头行，**其正上方必须是区块标题行**。
    没有这道闸门，单区块表里形态巧合的数据行会被误判成第二个表头，
    把一张表切成碎块、任务大面积丢失（实测 fixture_dirty 从 12 条掉到 5 条）。
    """
    heads = []
    for r in range(len(matrix)):
        nxt = matrix[r + 1] if r + 1 < len(matrix) else None
        if not _looks_like_block_header(matrix[r], nxt):
            continue
        if heads:
            prev = matrix[r - 1] if r > 0 else None
            if prev is None or not _is_block_title_row(prev):
                continue
        heads.append(r)
    blocks = []
    for i, r in enumerate(heads):
        end = (heads[i + 1] - 1) if i + 1 < len(heads) else len(matrix) - 1
        blocks.append({"header_row": r, "end": end})
    return blocks


SUBTOTAL_RE = re.compile(
    r"^(合\s*计|小\s*计|总\s*计|累\s*计|汇\s*总|总\s*额|平\s*均|均\s*值|本页合计|"
    r"total|sum|subtotal|average|avg)$", re.I)
NOTE_RE = re.compile(
    r"^(备\s*注|注[:：]?|说\s*明|温馨提示|填写说明|制表人|审核人|批准人|批准|"
    r"数据来源|编制|以下无正文)")
SEP_RE = re.compile(r"^[-—=_*·\.\s]{2,}$")


def classify_rows(matrix, header, headers) -> dict:
    """行分类：DATA / SUBTOTAL / NOTE / BLANK / SEPARATOR / HEADER_REPEAT。

    合计行**禁止用"数值特别大"判定**；无关键词时靠可加性自检兜底。
    """
    start = header["data_start"]
    labels = {}
    header_row = matrix[header["row_index"]]
    header_texts = [norm_str(v) for v in header_row if not is_null(v)]
    n_cols = len(headers)

    # 数值列索引（用于可加性自检）
    num_cols = []
    for c in range(n_cols):
        vals = [matrix[r][c] for r in range(start, min(start + 30, len(matrix)))]
        nums = [to_float(v) for v in vals if not is_null(v)]
        if nums and len(nums) / max(len([v for v in vals if not is_null(v)]), 1) >= 0.5:
            num_cols.append(c)

    for r in range(start, len(matrix)):
        row = matrix[r]
        nonempty = [(c, v) for c, v in enumerate(row) if not is_null(v)]
        if not nonempty:
            labels[r] = {"label": "BLANK", "reason": "blank"}
            continue

        first_text = norm_str(nonempty[0][1]) if isinstance(nonempty[0][1], str) else ""
        joined = " ".join(norm_str(v) for _, v in nonempty[:4])

        if first_text and SUBTOTAL_RE.match(first_text.replace(" ", "")):
            labels[r] = {"label": "SUBTOTAL", "reason": f"subtotal_keyword:{first_text}"}
            continue
        if NOTE_RE.match(first_text or joined):
            labels[r] = {"label": "NOTE", "reason": f"note_keyword:{first_text or joined[:10]}"}
            continue
        if len(nonempty) <= 1 and SEP_RE.match(first_text or norm_str(nonempty[0][1])):
            labels[r] = {"label": "SEPARATOR", "reason": "separator"}
            continue

        # 重复表头（分页打印）
        row_texts = [norm_str(v) for _, v in nonempty]
        if header_texts and len(row_texts) >= 2:
            common = sum(1 for t in row_texts if t in header_texts)
            if common / len(row_texts) >= P.DUP_HEADER_SIM:
                labels[r] = {"label": "HEADER_REPEAT", "reason": "duplicate_header"}
                continue

        # 可加性自检：本行 ≈ 上方连续若干行之和
        suspect = None
        for c in num_cols:
            v = to_float(row[c]) if c < len(row) else None
            if v is None or abs(v) < 1e-9:
                continue
            s = 0.0
            cnt = 0
            for j in range(r - 1, start - 1, -1):
                if labels.get(j, {}).get("label") == "SUBTOTAL":
                    break
                pv = to_float(matrix[j][c]) if c < len(matrix[j]) else None
                if pv is None:
                    continue
                s += pv
                cnt += 1
                if abs(v - s) <= max(1e-6, 0.01 * abs(v)) and cnt >= 2:
                    suspect = {"col": c, "sum": s, "span": cnt}
                    break
                if s > abs(v) * 1.02:
                    break
            if suspect:
                break

        if suspect:
            # 首列为空或含合计词 → 确认为合计行；否则只标记可疑，不剔除
            if not first_text:
                labels[r] = {"label": "SUBTOTAL",
                             "reason": f"additivity_check:≈上方{suspect['span']}行之和"}
            else:
                labels[r] = {"label": "DATA", "reason": "",
                             "suspicious_total": True}
            continue

        labels[r] = {"label": "DATA", "reason": ""}

    return labels


def build_headers(matrix, header: dict) -> list:
    """生成列名列表（双层则拼接 父|子）。"""
    r = header["row_index"]
    parent_row = matrix[r] if header["row_span"] == 2 else None
    child_row = matrix[r + 1] if header["row_span"] == 2 else matrix[r]
    n_cols = max((len(row) for row in matrix[:r + 2]), default=0)

    headers = []
    for c in range(n_cols):
        child = norm_str(child_row[c]) if c < len(child_row) else ""
        if header["row_span"] == 2 and parent_row is not None:
            parent = norm_str(parent_row[c]) if c < len(parent_row) else ""
            name = f"{parent}|{child}" if parent and parent != child else (child or parent)
        else:
            name = child
        headers.append(name)
    return headers
