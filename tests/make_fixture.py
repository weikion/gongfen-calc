"""生成一份「故意做脏」的测试 xlsx，用于验证解析层关键路径。

覆盖：
- 表头在第 3 行（上方有大标题 + 编制人行）
- 纵向合并单元格（任务名跨 2 行）
- 无「合计」二字但等于上方之和的伪合计行
- 工时混排 2人天 / 0.5 / 8h / 约3天 / 待定
- 金额混排 ¥12,000 / 1.2万 / 空
- 空值写法 "-"
- 优先级 P0–P3
- 声称 480h 但工期只有 2 天的矛盾行
- 真实的 160h 长周期任务（应与矛盾行区分：豁免打折）
- 真正的「合计」行、空行、备注行
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent / "fixture_dirty.xlsx"

HEADER = ["任务名称", "难度等级", "工时", "金额", "优先级",
          "开始日期", "结束日期", "协作方", "备注"]

ROWS = [
    ["A系统迁移一期", "高", 160, "¥120,000", "P0", "2026-01-05", "2026-03-20", "研发、测试、运维", ""],
    ["A系统迁移二期", "中", 40, "¥60,000", "P1", "2026-01-10", "2026-03-25", "研发、测试", ""],
    ["A系统迁移三期", "中", 24, "¥30,000", "P2", "2026-02-01", "2026-03-30", "测试", ""],
    ["数据看板重构", "高", 30, "1.2万", "P1", "2026-01-08", "2026-02-28", "研发、产品", ""],   # r7 纵向合并 r7:r8
    [None, "高", 15, "¥8,000", "P1", "2026-01-20", "2026-02-20", "研发", ""],                  # r8 合并区第二行
    ["接口联调", "中", 16, "¥8,000", "P2", "2026-01-15", "2026-02-15", "研发、测试", ""],
    ["需求调研", "高", "约3天", "¥15,000", "P1", "2026-01-05", "2026-01-20", "产品", ""],
    ["日常运维支持", "低", "8h", "-", "P3", "2026-01-05", "2026-01-31", "运维", ""],
    ["故障应急处理", "高", 480, "¥50,000", "P0", "2026-01-05", "2026-01-06", "全部门", "工期仅2天"],
    ["客户培训", "中", "待定", "¥5,000", "P3", "2026-02-01", "2026-02-05", "市场", ""],
    ["报表开发", "中", "2人天", "¥12,000", "P2", "2026-01-10", "2026-02-10", "数据", ""],
    ["性能优化", "高", 0.5, None, "P2", "2026-01-20", "2026-02-20", "研发", ""],
    ["合规申报", "低", 4, "¥3,000", "P0", "2026-01-06", "2026-01-15", "法务", ""],
]


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"

    ws["A1"] = "2026年Q1 研发中心任务清单（工时填报）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:I1")
    ws["A2"] = "编制人：张三　　审核人：李四　　编制日期：2026-01-05"
    ws.merge_cells("A2:I2")

    for c, h in enumerate(HEADER, 1):
        ws.cell(3, c, h).font = Font(bold=True)

    for i, row in enumerate(ROWS):
        r = 4 + i
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)

    ws.merge_cells(start_row=7, start_column=1, end_row=8, end_column=1)

    # 伪合计行：无「合计」二字，数值 = 上方可转数字单元格之和
    total_row = 4 + len(ROWS)
    hours_sum = 0.0
    money_sum = 0.0
    for i, row in enumerate(ROWS):
        hv, mv = row[2], row[3]
        if isinstance(hv, (int, float)):
            hours_sum += float(hv)
        if isinstance(mv, (int, float)):
            money_sum += float(mv)
        else:
            s = str(mv or "").replace("¥", "").replace(",", "")
            try:
                money_sum += float(s)
            except ValueError:
                pass
    ws.cell(total_row, 3, round(hours_sum, 2))
    ws.cell(total_row, 4, round(money_sum, 2))

    ws.cell(total_row + 1, 1, None)                       # 空行
    ws.cell(total_row + 2, 1, "备注：以上工时不含评审与返工时间")   # 备注行
    ws.cell(total_row + 3, 1, "合计")                      # 真正的合计行
    ws.cell(total_row + 3, 3, 9999)

    wb.save(OUT)
    print(f"written: {OUT}")
    print(f"伪合计行 row={total_row}, 工时和={hours_sum}, 金额和={money_sum}")


if __name__ == "__main__":
    build()
