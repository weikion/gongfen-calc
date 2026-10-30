"""生成 4 份「会被用户真实遇到、错了又很难发现」的 fixture。

优先补这 4 项（主理人指定）：
  ① 状态列仲裁为 progress —— 「状态」列全是 已完成/进行中/未开始，
     必须**不当作质量结果列**。错了会让所有已完成任务拿到同一个质量分，E 维退化成常数。
  ② 容器父任务可加性自检 —— 父工时 ≈ Σ子工时 → 父不计分（否则父子重复计分）；
     另给一组「父工时 >> 子之和」的真实独立父任务，其子任务应按 SUBTASK_FACTOR 折价。
  ③ Excel 序列号日期（40000–60000）—— 国内表极常见，序列号不能被当成工时/金额。
  ④ 双层表头 + 单位行 —— 父行纵向合并、子行细分，第三行是「单位:人天」行。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent


# ---------------------------------------------------------------- ① 状态列
def build_status_progress():
    wb = Workbook()
    ws = wb.active
    ws.title = "任务清单"
    header = ["任务名称", "状态", "工时", "金额", "优先级", "难度等级"]
    rows = [
        ["门户改版前端开发", "已完成", 60, 80000, "P0", "高"],
        ["订单服务接口重构", "进行中", 40, 50000, "P1", "高"],
        ["数据字典梳理", "未开始", 12, 5000, "P3", "低"],
        ["对账脚本编写", "已完成", 24, 20000, "P2", "中"],
        ["客户回访纪要", "进行中", 6, 2000, "P3", "低"],
        ["季度运营报告", "已完成", 16, 12000, "P2", "中"],
    ]
    for c, h in enumerate(header, 1):
        ws.cell(1, c, h).font = Font(bold=True)
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            ws.cell(2 + i, c, v)
    p = HERE / "fixture_status_progress.xlsx"
    wb.save(p)
    print(f"written: {p}")


# ---------------------------------------------------------------- ② 容器父任务
def build_container():
    wb = Workbook()
    ws = wb.active
    ws.title = "WBS"
    header = ["任务名称", "父任务", "工时", "金额", "难度等级", "开始日期", "结束日期"]
    rows = [
        # 父工时 40 = 15 + 15 + 10 → 判为容器，父不计分
        ["平台化建设项目", None, 40, 300000, "高", "2026-01-05", "2026-03-31"],
        ["后端服务开发", "平台化建设项目", 15, 120000, "高", "2026-01-05", "2026-02-15"],
        ["前端页面开发", "平台化建设项目", 15, 100000, "中", "2026-01-10", "2026-02-20"],
        ["联调与测试", "平台化建设项目", 10, 80000, "中", "2026-02-20", "2026-03-10"],
        # 父工时 100 >> 子之和 6 → 父是独立任务，子任务折价 0.6
        ["重大故障处置专项", None, 100, 200000, "高", "2026-01-08", "2026-02-28"],
        ["故障复盘报告", "重大故障处置专项", 6, 8000, "中", "2026-01-20", "2026-01-25"],
        # 完全独立的平级任务
        ["合规申报", None, 12, 15000, "中", "2026-01-06", "2026-01-15"],
    ]
    for c, h in enumerate(header, 1):
        ws.cell(1, c, h).font = Font(bold=True)
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            ws.cell(2 + i, c, v)
    p = HERE / "fixture_container.xlsx"
    wb.save(p)
    print(f"written: {p}")


# ---------------------------------------------------------------- ③ 序列号日期
def build_serial_date():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    header = ["任务名称", "工时", "开始日期", "结束日期", "金额"]
    # Excel 序列号（1899-12-30 起算）：2026-01-05=46027，2026-03-20=46101
    rows = [
        ["A系统迁移一期", 160, 46027, 46101, 120000],
        ["A系统迁移二期", 40, 46032, 46106, 60000],
        ["数据看板重构", 45, 46030, 46071, 12000],
        ["接口联调", 16, 46037, 46068, 8000],
        ["需求调研", 24, 46027, 46042, 15000],
        ["报表开发", 16, 46032, 46063, 12000],
        ["合规申报", 4, 46028, 46037, 3000],
    ]
    for c, h in enumerate(header, 1):
        ws.cell(1, c, h).font = Font(bold=True)
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            ws.cell(2 + i, c, v)
    p = HERE / "fixture_serial_date.xlsx"
    wb.save(p)
    print(f"written: {p}")


# ---------------------------------------------------------------- ④ 双层表头 + 单位行
def build_double_header():
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    # 父表头（第 1 行，纵向/横向合并）
    ws["A1"] = "任务名称"
    ws["B1"] = "开始日期"
    ws["C1"] = "工时投入"
    ws["E1"] = "金额"
    ws["G1"] = "交付物类型"
    # 子表头（第 2 行）
    ws["C2"] = "投入工时"
    ws["D2"] = "复核工时"
    ws["E2"] = "预算金额"
    ws["F2"] = "实际金额"
    # 单位行（第 3 行）
    ws["C3"] = "单位:人天"
    ws["D3"] = "单位:人天"
    ws["E3"] = "单位:万元"
    ws["F3"] = "单位:万元"

    for cell in ("A1", "B1", "C1", "E1", "G1", "C2", "D2", "E2", "F2"):
        ws[cell].font = Font(bold=True)
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.merge_cells("C1:D1")
    ws.merge_cells("E1:F1")
    ws.merge_cells("G1:G2")

    rows = [
        ["智能客服知识库建设", 46027, 12, 4, 60, 55, "手册"],
        ["工单自动分派脚本", 46032, 6, 2, 20, 18, "脚本"],
        ["客户满意度调研", 46030, 3, 1, 8, 6, "报告"],
        ["运营周报模板", 46035, 1.5, 0.5, 2, 2, "模板"],
        ["供应商对账", 46040, 4, 2, 15, 12, "报告"],
    ]
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            ws.cell(4 + i, c, v)
    p = HERE / "fixture_double_header.xlsx"
    wb.save(p)
    print(f"written: {p}")


def build_all():
    build_status_progress()
    build_container()
    build_serial_date()
    build_double_header()


if __name__ == "__main__":
    build_all()
